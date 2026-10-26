"""Build a Project Journal .docx from its Markdown source, using Journal 1 as the layout template.

Markdown format (see docs/journal/TEMPLATE.md):
  # Project Journal Submission N (Covering week X-Y)     -> header line
  Key: value lines before the first "##"                -> header fields
  ## Section name                                       -> one table row, in template order
  **text**   -> bold line,   - text -> bullet,   anything else -> plain paragraph

Usage:  python tools/build_journal.py docs/journal/journal_02_weeks_4-5.md
"""
import copy
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "journal" / "Project Journal 1(Week 1-3).docx"
SECTIONS = ["Title of dissertation and brief description", "Communicating with supervisor",
            "References consulted", "Tools explored/used", "Other work carried out",
            "Plan the next 2 to 3 weeks", "Overall reflection", "Additional section, if needed"]


def parse(md: str):
    header, sections, current = [], {}, None
    for line in md.splitlines():
        if line.startswith("## "):
            current = line[3:].strip()
            sections[current] = []
        elif line.strip():
            (sections[current] if current else header).append(line.strip())
    return header, sections


def set_text(par, text, bold=None):
    runs = par.runs
    runs[0].text = text
    if bold is not None:
        runs[0].bold = bold
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def fill_cell(cell, lines, protos):
    for p in list(cell.paragraphs):
        p._p.getparent().remove(p._p)
    for line in lines:
        if line.startswith("**") and line.endswith("**"):
            kind, text = "bold", line[2:-2]
        elif line.startswith("- "):
            kind, text = "bullet", line[2:]
        else:
            kind, text = "plain", line
        el = copy.deepcopy(protos[kind])
        cell._tc.append(el)
        set_text(docx.text.paragraph.Paragraph(el, cell), text, bold=(kind == "bold"))


def build(md_path: Path) -> Path:
    header, sections = parse(md_path.read_text(encoding="utf-8"))
    d = docx.Document(TEMPLATE)
    rows = d.tables[0].rows
    protos = {"bold": copy.deepcopy(rows[2].cells[1].paragraphs[0]._p),
              "bullet": copy.deepcopy(rows[2].cells[1].paragraphs[1]._p),
              "plain": copy.deepcopy(rows[7].cells[1].paragraphs[0]._p)}

    title = header[0].lstrip("# ").strip()
    m = re.match(r"Project Journal Submission\s+(.*)", title)
    hp = rows[0].cells[0].paragraphs
    hp[0].runs[0].text = "Project Journal Submission "
    set_text_tail(hp[0], m.group(1) if m else title)
    fields = dict(l.split(":", 1) for l in header[1:] if ":" in l)
    for p in hp[1:]:
        key = p.text.split(":", 1)[0]
        if key in fields:
            set_text(p, f"{key}: {fields[key].strip()}")

    for i, name in enumerate(SECTIONS, start=1):
        if name in sections:
            fill_cell(rows[i].cells[1], sections[name], protos)

    out = md_path.with_suffix(".docx")
    d.save(out)
    return out


def set_text_tail(par, text):
    par.runs[1].text = text
    for r in par.runs[2:]:
        r._r.getparent().remove(r._r)


if __name__ == "__main__":
    print(build(Path(sys.argv[1])))
